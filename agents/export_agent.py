import base64
import os
import requests
from fpdf import FPDF
from docx import Document
from docx.shared import Inches, Pt
from PIL import Image as PILImage
import io


def _fetch_image_bytes(url: str) -> bytes:
    """Return raw image bytes from either an http(s) URL or a data: URI."""
    if url.startswith("data:"):
        _, payload = url.split(",", 1)
        payload = payload.strip()
        # Fix missing base64 padding
        missing = len(payload) % 4
        if missing:
            payload += "=" * (4 - missing)
        return base64.b64decode(payload)
    resp = requests.get(url, timeout=20)
    resp.raise_for_status()
    return resp.content


def _image_buf(url: str) -> io.BytesIO | None:
    """Return a BytesIO JPEG buffer, or None if anything fails (with logging)."""
    try:
        raw = _fetch_image_bytes(url)
        pil = PILImage.open(io.BytesIO(raw)).convert("RGB")
        buf = io.BytesIO()
        pil.save(buf, format="JPEG")
        buf.seek(0)
        return buf
    except Exception as exc:
        import traceback
        print(f"[WARN] Image skipped: {exc}")
        traceback.print_exc()
        return None

_UNICODE_MAP = {
    "‘": "'",  "’": "'",   # curly single quotes
    "“": '"',  "”": '"',   # curly double quotes
    "–": "-",  "—": "--",  # en-dash, em-dash
    "…": "...",                  # ellipsis
    " ": " ",                    # non-breaking space
    "•": "-",  "‣": "-",   # bullet variants
    "®": "(R)", "©": "(C)", # (R) (C)
    "¼": "1/4", "½": "1/2", "¾": "3/4",
}

def _to_latin1(text: str) -> str:
    for ch, rep in _UNICODE_MAP.items():
        text = text.replace(ch, rep)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def export_to_pdf(topic: str, content: str, image_url: str = None) -> str:
    """
    Export blog post to PDF.
    ✅ Compatible with fpdf2 2.8.x and Pillow 12.x
    """
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    NX, NY = "LMARGIN", "NEXT"

    # Title
    pdf.set_font("Helvetica", "B", 22)
    pdf.set_text_color(30, 30, 60)
    pdf.multi_cell(0, 12, _to_latin1(topic), new_x=NX, new_y=NY)
    pdf.ln(6)

    # Divider
    pdf.set_draw_color(123, 97, 255)
    pdf.set_line_width(0.8)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(8)

    # Cover image
    if image_url:
        print("[PDF] Loading cover image…")
        buf = _image_buf(image_url)
        if buf:
            pdf.image(buf, x=10, w=190)
            pdf.ln(8)

    # Content — parse markdown headings
    pdf.set_font("Helvetica", size=11)
    pdf.set_text_color(40, 40, 40)

    for line in content.split("\n"):
        line = line.strip()
        if not line:
            pdf.ln(4)
            continue

        if line.startswith("## "):
            pdf.set_font("Helvetica", "B", 14)
            pdf.set_text_color(123, 97, 255)
            pdf.multi_cell(0, 8, _to_latin1(line[3:]), new_x=NX, new_y=NY)
            pdf.set_font("Helvetica", size=11)
            pdf.set_text_color(40, 40, 40)
            pdf.ln(2)
        elif line.startswith("# "):
            pdf.set_font("Helvetica", "B", 18)
            pdf.set_text_color(30, 30, 60)
            pdf.multi_cell(0, 10, _to_latin1(line[2:]), new_x=NX, new_y=NY)
            pdf.set_font("Helvetica", size=11)
            pdf.set_text_color(40, 40, 40)
            pdf.ln(3)
        elif line.startswith("- "):
            pdf.multi_cell(0, 7, _to_latin1(f"  - {line[2:]}"), new_x=NX, new_y=NY)
        else:
            pdf.multi_cell(0, 7, _to_latin1(line), new_x=NX, new_y=NY)

    # Save
    safe = "".join(c for c in topic if c.isalnum() or c == " ")[:25]
    output_path = f"blog_{safe.replace(' ', '_')}.pdf"
    pdf.output(output_path)
    print(f"[DONE] PDF saved: {output_path}")
    return output_path


def export_to_word(topic: str, content: str, image_url: str = None) -> str:
    """
    Export blog post to Word (.docx).
    ✅ Compatible with python-docx 1.2.x
    """
    doc = Document()

    # Title
    title_para = doc.add_heading(topic, level=0)
    title_para.runs[0].font.size = Pt(24)
    doc.add_paragraph()

    # Cover image
    if image_url:
        print("[WORD] Loading cover image…")
        buf = _image_buf(image_url)
        if buf:
            doc.add_picture(buf, width=Inches(6))
            doc.add_paragraph()

    # Content — parse markdown headings
    for line in content.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("## "):
            doc.add_heading(line.replace("## ", ""), level=2)
        elif line.startswith("# "):
            doc.add_heading(line.replace("# ", ""), level=1)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)

    # Save
    safe = "".join(c for c in topic if c.isalnum() or c == " ")[:25]
    output_path = f"blog_{safe.replace(' ', '_')}.docx"
    doc.save(output_path)
    print(f"[DONE] Word doc saved: {output_path}")
    return output_path
